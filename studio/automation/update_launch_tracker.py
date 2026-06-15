from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

ROOT = Path(".").resolve()
TRACKER = ROOT / "truth_explains_launch_tracker.xlsx"
EPISODES_DIR = ROOT / "episodes"

if not TRACKER.exists():
    raise FileNotFoundError("truth_explains_launch_tracker.xlsx not found in project root")

wb = load_workbook(TRACKER)

def get_or_create_sheet(name, headers):
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
        ws.append(headers)
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        ws.append(headers)
    return ws

episodes_ws = get_or_create_sheet("Episode Index", [
    "Episode",
    "Folder",
    "Script",
    "Beats",
    "Prompts",
    "Still Images",
    "Runway Clips",
    "Voiceover",
    "Music",
    "SFX",
    "Resolve",
    "Export",
    "Status",
    "Last Updated"
])

automation_ws = get_or_create_sheet("Automation Log", [
    "Timestamp",
    "Action",
    "Details"
])

existing_rows = {}
for row in range(2, episodes_ws.max_row + 1):
    episode_name = episodes_ws.cell(row, 1).value
    if episode_name:
        existing_rows[episode_name] = row

def count_files(path, extensions=None):
    if not path.exists():
        return 0
    files = [p for p in path.iterdir() if p.is_file()]
    if extensions:
        files = [p for p in files if p.suffix.lower() in extensions]
    return len(files)

def exists_status(path):
    return "Yes" if path.exists() and count_files(path) > 0 else "No"

for ep_folder in sorted(EPISODES_DIR.iterdir()):
    if not ep_folder.is_dir():
        continue

    episode_name = ep_folder.name

    script = exists_status(ep_folder / "01_script")
    beats = exists_status(ep_folder / "02_beats")
    prompts = exists_status(ep_folder / "03_prompts")
    stills = count_files(ep_folder / "04_still_images", [".png", ".jpg", ".jpeg", ".webp"])
    clips = count_files(ep_folder / "05_runway_clips", [".mp4", ".mov"])
    voiceover = count_files(ep_folder / "06_voiceover", [".wav", ".mp3", ".aiff"])
    music = count_files(ep_folder / "07_music", [".wav", ".mp3", ".aiff"])
    sfx = count_files(ep_folder / "08_sfx", [".wav", ".mp3", ".aiff"])
    resolve = count_files(ep_folder / "09_resolve", [".drp", ".dra"])
    export = count_files(ep_folder / "10_export", [".mp4", ".mov"])

    if export > 0:
        status = "Live or Exported"
    elif clips > 0:
        status = "In Edit"
    elif stills > 0:
        status = "In Production"
    elif script == "Yes" or beats == "Yes":
        status = "In Development"
    else:
        status = "Created"

    row_data = [
        episode_name,
        str(ep_folder),
        script,
        beats,
        prompts,
        stills,
        clips,
        voiceover,
        music,
        sfx,
        resolve,
        export,
        status,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ]

    if episode_name in existing_rows:
        row = existing_rows[episode_name]
        for col, value in enumerate(row_data, start=1):
            episodes_ws.cell(row, col).value = value
    else:
        episodes_ws.append(row_data)

automation_ws.append([
    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    "Launch tracker updated",
    "Scanned episodes folder and refreshed Episode Index"
])

wb.save(TRACKER)
print("Launch tracker updated successfully")
